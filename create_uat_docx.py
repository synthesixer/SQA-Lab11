import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# สร้างเอกสารใหม่
doc = docx.Document()

# ตั้งหัวเรื่องหลัก
title = doc.add_heading('UAT Test Document – ระบบประกันรถยนต์', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# เพิ่มส่วนข้อมูลการทดสอบ
p = doc.add_paragraph()
info_table = doc.add_table(rows=5, cols=2)
info_table.style = 'Table Grid'
info_data = [
    ('ชื่อไฟล์ Test Script', 'Lab11_UAT_TestScript.docx'),
    ('วันที่ทดสอบ', '2026‑09‑14'),
    ('ผู้ทดสอบ', 'QA Engineer (อัตโนมัติ)'),
    ('สภาพแวดล้อม', 'Windows 11, Chrome 118, Tricentis Tosca / Selenium'),
    ('เครื่องมืออัตโนมัติ', 'python‑docx')
]
for i, (k, v) in enumerate(info_data):
    row = info_table.rows[i]
    row.cells[0].text = k
    row.cells[1].text = v
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.name = 'TH SarabunPSK'
                run.font.size = Pt(14)

# เพิ่มตาราง Test Scenario
doc.add_paragraph()
heading = doc.add_heading('ผลการทดสอบ (Test Scenario)', level=1)
heading.runs[0].font.name = 'TH SarabunPSK'
heading.runs[0].font.size = Pt(16)

# สร้างตารางหลัก
table = doc.add_table(rows=1, cols=4)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_titles = ['Test Scenario', 'Actual Result', 'Test Result (Pass/Fail)', 'Remark']
for i, title in enumerate(hdr_titles):
    hdr_cells[i].text = title
    hdr_par = hdr_cells[i].paragraphs[0]
    hdr_par.runs[0].font.name = 'TH SarabunPSK'
    hdr_par.runs[0].font.size = Pt(14)
    hdr_par.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Test Case 1 – ขอใบเสนอราคาสำเร็จ
row_cells = table.add_row().cells
row_cells[0].text = 'TC01 – ขอใบเสนอราคาสำเร็จ (ข้อมูลถูกต้อง)'
row_cells[1].text = 'ระบบแสดงฟอร์มครบถ้วน ส่งอีเมลสำเร็จและแสดงข้อความ "ส่งสำเร็จ"'
row_cells[2].text = 'Pass'
row_cells[3].text = ''

# Test Case 2 – ขอใบเสนอราคาไม่สำเร็จ (ข้อมูลไม่ครบ/ผิดพลาด)
row_cells = table.add_row().cells
row_cells[0].text = 'TC02 – ขอใบเสนอราคาไม่สำเร็จ (ข้อมูลไม่ครบ/ผิดพลาด)'
row_cells[1].text = 'ระบบแสดงข้อผิดพลาด "กรุณากรอกข้อมูลรถให้ครบถ้วน" ไม่ส่งอีเมล'
row_cells[2].text = 'Pass'
row_cells[3].text = ''

# ตั้งฟอนต์ให้ทุกเซลล์
for row in table.rows:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.name = 'TH SarabunPSK'
                run.font.size = Pt(14)

# บันทึกไฟล์
output_path = r'e:/Doc/Code/GitHub/SQA/SQALAB11/SQA-Lab11/Docs/Lab11_UAT_Document.docx'
doc.save(output_path)
print('DOCX created at', output_path)
